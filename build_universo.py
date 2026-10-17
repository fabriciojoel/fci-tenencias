#!/usr/bin/env python3
"""universo_sitio.json = [[nombre, ticker D, vencimiento ISO|None]] desde el Excel del modelo de ON."""
import datetime
import json
import os
import openpyxl

XL = "E:/Inversiones Fabri-Claude/ON BYMA/ON u$s Automatica - modelo con formulas.xlsx"
wb = openpyxl.load_workbook(XL, read_only=True, data_only=True)
tabs = set(wb.sheetnames)
out = []
for r in wb["ONs"].iter_rows(min_row=4, values_only=True):
    if not (r[0] and isinstance(r[1], str) and len(r[1]) == 5):
        continue
    venc = None
    for t in (r[2], r[1]):
        if t in tabs:
            fechas = [row[0] for row in wb[t].iter_rows(min_row=20, max_col=1, values_only=True)
                      if isinstance(row[0], datetime.datetime)]
            if fechas:
                venc = max(fechas).date().isoformat()
            break
    out.append([r[0], r[1], venc])
json.dump(out, open(os.path.join(os.path.dirname(__file__), "universo_sitio.json"), "w", encoding="utf-8"),
          ensure_ascii=False)
print(len(out), "con vencimiento:", sum(1 for x in out if x[2]))
